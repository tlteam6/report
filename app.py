import streamlit as st
import pdfplumber
import openpyxl
import re
import json
from collections import defaultdict
import gspread
from google.oauth2.service_account import Credentials

# ---------------------------------------------------------------------------
# ตั้งค่าหน้าเว็บ
# ---------------------------------------------------------------------------
st.set_page_config(page_title="Import รายงานเบี้ยประกัน", page_icon="📄")
st.title("📄 Import รายงานเบี้ยประกัน")
st.write("Upload รายงานเบี้ยประกัน PDF ระบบจะอ่านและเขียนข้อมูลเข้า Google Sheet ให้อัตโนมัติ")

# ---------------------------------------------------------------------------
# ตรรกะแปลง PDF (เหมือนที่ทดสอบไว้แล้ว)
# ---------------------------------------------------------------------------
COLS = [
    ("ค่าบำเหน็จ", 225, 300),
    ("เบี้ยเครดิต", 300, 370),
    ("ปีต่อไป", 370, 432),
    ("รวม", 432, 496),
    ("ฐานเบี้ย", 496, 540),
    ("%เบี้ยปีแรก/ฐาน", 540, 999),
]
NUM_RE = re.compile(r"^-?[\d,]+(\.\d+)?$")
SKIP_SUBSTRINGS = [
    "วันที่ออกรายงาน", "รายงานเบี้ยประกัน", "ข้อมูลเบี้ยตัด",
    "หมายเหตุ", "ตัวเลขด้านล่าง", "ที่มา", "โครงการนักขายดิจิทัล",
    "เบี้ยประกันรับ", "ฐานเบี้ย", "ค่าบำเหน็จ",
]
HEADER_WORDS = {"ปีแรก", "ปีต่อไป", "รวม", "แรก/ฐาน", "%", "เบี้ยปี", "ตำแหน่ง", "ชื่อ", "-", "สกุล"}


def is_number(tok):
    return bool(NUM_RE.match(tok))


def col_for_x(x0):
    for name, lo, hi in COLS:
        if lo <= x0 < hi:
            return name
    return None


def extract_lines(file_obj):
    all_lines = []
    with pdfplumber.open(file_obj) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            words = sorted(page.extract_words(), key=lambda w: w["top"])
            groups = []
            current_group = []
            current_top = None
            for w in words:
                if current_top is not None and abs(w["top"] - current_top) <= 3:
                    current_group.append(w)
                else:
                    if current_group:
                        groups.append(current_group)
                    current_group = [w]
                    current_top = w["top"]
            if current_group:
                groups.append(current_group)
            for g in groups:
                ws = sorted(g, key=lambda w: w["x0"])
                all_lines.append((pno, g[0]["top"], ws))
    return all_lines


def should_skip(ws):
    texts = [w["text"] for w in ws]
    joined = " ".join(texts)
    if any(s in joined for s in SKIP_SUBSTRINGS):
        return True
    if set(texts) <= HEADER_WORDS:
        return True
    return False


def is_continuation_number_line(ws):
    texts = [w["text"] for w in ws]
    if not texts:
        return False
    return all(is_number(t) for t in texts)


def parse_data_line(ws, current_label):
    texts_with_x = [(w["text"], w["x0"]) for w in ws]
    idx = 0
    label = current_label
    if texts_with_x and texts_with_x[0][1] < 40:
        label = texts_with_x[0][0]
        idx = 1
    code = ""
    if idx < len(texts_with_x) and texts_with_x[idx][1] < 95:
        code = texts_with_x[idx][0]
        idx += 1
    name_tokens = []
    while idx < len(texts_with_x) and texts_with_x[idx][1] < 225:
        name_tokens.append(texts_with_x[idx][0])
        idx += 1
    name = " ".join(name_tokens).replace(" - ", "-").strip()

    if not code and name:
        first_word, _, rest = name.partition(" ")
        if re.match(r"^[\d.\-/]+$", first_word):
            code = first_word
            name = rest.strip()

    if re.match(r"^\d{5}$", code):
        label = "ตัวแทน"

    row = {
        "ตำแหน่ง": label, "รหัส": code, "ชื่อ-สกุล": name,
        "ค่าบำเหน็จ": "", "เบี้ยเครดิต": "", "เบี้ยไม่คิดผลงาน": "",
        "ปีต่อไป": "", "รวม": "", "ฐานเบี้ย": "", "%เบี้ยปีแรก/ฐาน": "",
    }
    for text, x0 in texts_with_x[idx:]:
        col = col_for_x(x0)
        if col:
            row[col] = text
    return row


def parse_pdf(file_obj):
    lines = extract_lines(file_obj)
    rows = []
    current_label = ""
    current_fah = ""
    current_pak = ""
    current_soon = ""
    current_nuay = ""

    for pno, top, ws in lines:
        if should_skip(ws):
            continue
        if is_continuation_number_line(ws):
            if rows:
                rows[-1]["เบี้ยไม่คิดผลงาน"] = ws[0]["text"]
            continue

        row = parse_data_line(ws, current_label)
        current_label = row["ตำแหน่ง"]
        name = row["ชื่อ-สกุล"]

        if current_label == "ฝ่าย":
            current_fah = name
            current_pak = ""
            current_soon = ""
            current_nuay = ""
        elif current_label in ("ภาค", "ภาค(GL)"):
            current_pak = name
            current_soon = ""
            current_nuay = ""
        elif current_label == "ศูนย์":
            current_soon = name
            current_nuay = ""
        elif current_label == "หน่วย":
            current_nuay = name

        row["ฝ่าย"] = current_fah
        row["ภาค"] = current_pak
        row["ศูนย์"] = current_soon
        row["หน่วย"] = current_nuay

        rows.append(row)
    return rows


FIELDNAMES = ["ตำแหน่ง", "รหัส", "ชื่อ-สกุล", "ฝ่าย", "ภาค", "ศูนย์", "หน่วย",
              "ค่าบำเหน็จ", "เบี้ยเครดิต", "เบี้ยไม่คิดผลงาน", "ปีต่อไป", "รวม",
              "ฐานเบี้ย", "%เบี้ยปีแรก/ฐาน"]


VALID_POSITION_LABELS = {"ฝ่าย", "ภาค", "ศูนย์", "หน่วย", "ตัวแทน"}


def parse_xlsx_col_a(text):
    left, _, name = text.partition(":")
    left = left.strip()
    name = name.strip()
    parts = left.split(None, 1)
    label = parts[0] if parts else ""
    code_raw = parts[1].strip() if len(parts) > 1 else ""
    middle = code_raw
    if middle and middle[0].isalpha():
        middle = middle[1:]
    if middle and middle[-1].isalpha():
        middle = middle[:-1]
    return label, middle, name


def match_code_for_label(label, code):
    if label == "ตัวแทน":
        return code[-5:]
    return code


def parse_xlsx(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    result = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        col_a, col_c = row[0], row[2]
        if not col_a:
            continue
        label, code, name = parse_xlsx_col_a(str(col_a))
        if label not in VALID_POSITION_LABELS:
            # ข้ามแถวท้ายไฟล์ที่ไม่ใช่ข้อมูลจริง เช่น "Applied filters: ..."
            continue
        match_code = match_code_for_label(label, code)
        result.append((label, match_code, col_c, strip_title_prefix(name)))
    return result


FYP_MONTHS = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
              "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
FYP_HEADER = ["ชื่อ", "ตำแหน่ง", "รหัส"] + FYP_MONTHS + ["สะสม"]

TITLE_PREFIX_RE = re.compile(r"^(นางสาว|นาย|นาง|น\.ส\.)\s*")


def strip_title_prefix(name):
    """ตัดคำนำหน้าชื่อออก เช่น 'นายสุภัทรชัย โกศาคาร' -> 'สุภัทรชัย โกศาคาร'"""
    return TITLE_PREFIX_RE.sub("", name).strip()


def find_fyp_month_columns(ws):
    """หาว่าคอลัมน์ไหนในไฟล์ต้นทางคือเดือนอะไร จากหัวตาราง เช่น 'ม.ค. 69' -> 'ม.ค.'
    คืนค่า dict: column_index(1-based) -> ชื่อเดือนแบบย่อ"""
    month_col_map = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=c).value
        if not val or not isinstance(val, str):
            continue
        token = val.strip().split()[0] if val.strip() else ""
        if token in FYP_MONTHS:
            month_col_map[c] = token
    return month_col_map


def parse_fyp_xlsx(file_obj):
    """อ่านไฟล์ xlsx แบบรายเดือน (คอลัมน์ A=ตำแหน่ง+รหัส+ชื่อ, ตามด้วยคอลัมน์ค่าต่อเดือน)
    คืนค่า list of dict: {label, code, name, months: {ชื่อเดือน: ค่า}}"""
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    month_col_map = find_fyp_month_columns(ws)

    result = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        col_a = row[0]
        if not col_a:
            continue
        label, code, name = parse_xlsx_col_a(str(col_a))
        if label not in VALID_POSITION_LABELS:
            # ข้ามแถวท้ายไฟล์ที่ไม่ใช่ข้อมูลจริง เช่น "Applied filters: ..."
            continue
        clean_name = strip_title_prefix(name)
        months = {}
        for col_idx, month in month_col_map.items():
            val = row[col_idx - 1] if col_idx - 1 < len(row) else None
            if val not in (None, ""):
                months[month] = val
        result.append({"label": label, "code": code, "name": clean_name, "months": months})
    return result


CLOSEBOOK_HEADER = ["ชื่อ-สกุล", "ตำแหน่ง", "ภาค", "ศูนย์", "หน่วย",
                     "เบี้ยปีแรก", "ค่าบำเหน็จ", "%เบี้ยปีแรก", "เบี้ยต่อไป"]


def parse_closebook_xlsx(file_obj):
    """อ่านไฟล์ xlsx วันปิดบัญชี (หัวตารางแถวที่ 1, ข้อมูลเริ่มแถวที่ 2)
    คืนค่า list of [ชื่อ, ตำแหน่ง, เบี้ยปีแรก, ค่าบำเหน็จ, %เบี้ยปีแรก, เบี้ยต่อไป]"""
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    result = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        col_a = row[0]
        if not col_a:
            continue
        label, code, name = parse_xlsx_col_a(str(col_a))
        if label not in VALID_POSITION_LABELS:
            # ข้ามแถวท้ายไฟล์ที่ไม่ใช่ข้อมูลจริง เช่น "Applied filters: ..."
            continue
        clean_name = strip_title_prefix(name)
        first_year = row[2] if len(row) > 2 else None
        commission = row[3] if len(row) > 3 else None
        first_year_pct = row[4] if len(row) > 4 else None
        renewal = row[5] if len(row) > 5 else None
        result.append([clean_name, label, first_year, commission, first_year_pct, renewal])
    return result


def extract_date_from_filename(filename):
    match = re.search(r"(\d{2})-(\d{2})-(\d{2})", filename)
    if not match:
        return None
    dd, mm, yy = match.groups()
    return f"{dd}/{mm}/25{yy}"


def get_gsheet_client():
    scopes = ["https://www.googleapis.com/auth/spreadsheets",
              "https://www.googleapis
